"""HR Library service — upload, text extraction, CRUD."""

from __future__ import annotations

import logging
import os
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError

from biotact.modules.hr.library.models import HRTemplate
from biotact.modules.hr.library.scanner import scan_template_fields

if TYPE_CHECKING:
    from fastapi import UploadFile
    from sqlalchemy.ext.asyncio import AsyncSession
from biotact.core.config import get_settings
from biotact.modules.hr.library.schemas import (
    TemplateDetailResponse,
    TemplateListResponse,
    TemplateResponse,
)

logger = logging.getLogger(__name__)

ALLOWED_EXTS = frozenset({"docx", "pdf", "txt", "md"})
# Magic bytes for content-type verification (text formats have no fixed signature)
_MAGIC_BYTES: dict[str, bytes] = {
    "docx": b"PK\x03\x04",  # DOCX is a ZIP archive
    "pdf": b"%PDF-",
}


class HRFileError(ValueError):
    """Base class for HR template upload validation errors."""

    status_code = 400


class HRFileTypeError(HRFileError):
    """Unsupported file extension."""


class HRFileMagicError(HRFileError):
    """Declared file type does not match content magic bytes."""


class HRFileTooLarge(HRFileError):
    """Uploaded file exceeds the size limit."""

    status_code = 413


class HRTemplateConflictError(HRFileError):
    """Concurrent upload produced a version conflict."""

    status_code = 409


def _upload_dir() -> Path:
    """Return configured upload directory."""
    return Path(get_settings().hr_upload_dir)


def _ensure_upload_dir() -> None:
    """Create upload directory if it doesn't exist."""
    _upload_dir().mkdir(parents=True, exist_ok=True)


def _verify_magic_bytes(ext: str, first_chunk: bytes) -> None:
    """Raise HRFileMagicError if declared ext doesn't match content header.

    Text formats (``txt``, ``md``) and any extension absent from
    ``_MAGIC_BYTES`` skip the check — they have no fixed signature.
    """
    expected_magic = _MAGIC_BYTES.get(ext)
    if expected_magic and not first_chunk.startswith(expected_magic):
        msg = f"File content does not match declared type '{ext}'."
        raise HRFileMagicError(msg)


def _extract_text_docx(file_path: str) -> str:
    """Extract full text from DOCX file (paragraphs + table cells)."""
    doc = Document(file_path)
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        parts.append(p.text)
    return "\n".join(parts)


def _extract_text_pdf(file_path: str) -> str:
    """Extract full text from PDF file."""
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_text_txt(file_path: str) -> str:
    """Read plain text file."""
    with open(file_path, encoding="utf-8") as f:
        return f.read()


def extract_text(file_path: str, file_type: str) -> str | None:
    """Extract text from uploaded file based on type."""
    extractors = {
        "docx": _extract_text_docx,
        "pdf": _extract_text_pdf,
        "txt": _extract_text_txt,
        "md": _extract_text_txt,
    }
    extractor = extractors.get(file_type)
    if not extractor:
        logger.warning("No text extractor for file type: %s", file_type)
        return None
    try:
        return extractor(file_path)
    except Exception:
        logger.exception("Failed to extract text from %s", file_path)
        return None


async def upload_template(
    db: AsyncSession,
    file: UploadFile,
    category: str,
    user_id: int,
) -> TemplateResponse:
    """Upload a template file, extract text, save to DB.

    Validates filename (strips path components), file type, content magic
    bytes (where applicable), and streams write with a hard size limit.

    Raises:
        HRFileTypeError: extension not in ALLOWED_EXTS.
        HRFileMagicError: declared type doesn't match content header.
        HRFileTooLarge: upload exceeds MAX_UPLOAD_BYTES.
    """
    _ensure_upload_dir()

    settings = get_settings()
    max_upload_bytes = settings.hr_max_upload_mb * 1024 * 1024

    # Sanitize filename — strip any path components (cross-platform path traversal guard)
    raw_name = file.filename or "unknown"
    safe_basename = os.path.basename(raw_name.replace("\\", "/"))
    ext = safe_basename.rsplit(".", 1)[-1].lower() if "." in safe_basename else ""
    if ext not in ALLOWED_EXTS:
        msg = f"Unsupported file type: {ext or '<none>'}. Use docx, pdf, txt, or md."
        raise HRFileTypeError(msg)

    # Storage uses only uuid + ext — attacker-controlled bytes never on disk path
    storage_name = f"{uuid.uuid4().hex}.{ext}"
    file_path = _upload_dir() / storage_name

    # Verify magic bytes BEFORE persisting anything (no disk write on type mismatch)
    first_chunk = await file.read(8)
    _verify_magic_bytes(ext, first_chunk)

    # Stream the upload to disk, then extract text and persist DB row.
    # Any failure in this block (IOError, oversize, DB error, etc.) unlinks
    # the partial file — guards against orphan files on disk.
    total_size = len(first_chunk)
    try:
        with file_path.open("wb") as f:
            f.write(first_chunk)
            while chunk := await file.read(65536):
                total_size += len(chunk)
                if total_size > max_upload_bytes:
                    raise HRFileTooLarge(
                        f"Upload exceeds {settings.hr_max_upload_mb} MB limit."
                    )
                f.write(chunk)

        # Extract text
        extracted = extract_text(str(file_path), ext)

        # Scan for {{ PLACEHOLDER }} fields in DOCX templates
        fields: list[str] = []
        if ext == "docx":
            fields = scan_template_fields(str(file_path))

        # Versioning: lock current active template, bump version, deactivate prev.
        try:
            prev_result = await db.execute(
                select(HRTemplate)
                .where(HRTemplate.category == category)
                .where(HRTemplate.is_active.is_(True))
                .with_for_update()
            )
            prev = prev_result.scalar_one_or_none()
            next_version = (prev.version + 1) if prev else 1

            if prev is not None:
                prev.is_active = False
                await db.flush()

            template = HRTemplate(
                name=safe_basename,
                category=category,
                file_path=str(file_path),
                file_type=ext,
                file_size=total_size,
                extracted_text=extracted,
                template_fields=fields if fields else None,
                uploaded_by=user_id,
                version=next_version,
                is_active=True,
            )
            db.add(template)
            await db.flush()
            await db.refresh(template)

            if prev is not None:
                prev.superseded_by_id = template.id
                await db.flush()
        except IntegrityError as exc:
            file_path.unlink(missing_ok=True)
            msg = (
                f"Version conflict for category '{category}'. "
                "Another upload completed concurrently."
            )
            raise HRTemplateConflictError(msg) from exc
    except Exception:
        file_path.unlink(missing_ok=True)
        raise

    logger.info(
        "Template uploaded: id=%d name=%s category=%s fields=%s chars=%d size=%d",
        template.id,
        safe_basename,
        category,
        fields,
        len(extracted) if extracted else 0,
        total_size,
    )
    return TemplateResponse.model_validate(template)


async def list_templates(
    db: AsyncSession,
    category: str | None = None,
) -> TemplateListResponse:
    """List all templates, optionally filtered by category."""
    query = select(HRTemplate).order_by(HRTemplate.created_at.desc())
    if category:
        query = query.where(HRTemplate.category == category)

    result = await db.execute(query)
    items = list(result.scalars().all())

    count_query = select(func.count(HRTemplate.id))
    if category:
        count_query = count_query.where(HRTemplate.category == category)
    total = (await db.execute(count_query)).scalar_one()

    return TemplateListResponse(
        items=[TemplateResponse.model_validate(t) for t in items],
        total=total,
    )


async def get_template(
    db: AsyncSession,
    template_id: int,
) -> TemplateDetailResponse | None:
    """Get template with extracted text by ID."""
    result = await db.execute(select(HRTemplate).where(HRTemplate.id == template_id))
    template = result.scalar_one_or_none()
    if not template:
        return None
    return TemplateDetailResponse.model_validate(template)


async def get_template_by_category(
    db: AsyncSession,
    category: str,
) -> TemplateDetailResponse | None:
    """Get the active, render-eligible template for a category.

    Render-eligibility is gated on ``template_fields IS NOT NULL`` (the fields a
    document needs), NOT on ``extracted_text`` — extraction is unrelated to
    rendering (``chat/documents.py`` renders from ``file_path`` + fields), and in
    prod most templates have ``extracted_text = NULL`` yet render fine.
    """
    result = await db.execute(
        select(HRTemplate)
        .where(HRTemplate.category == category)
        .where(HRTemplate.is_active.is_(True))
        .where(HRTemplate.template_fields.is_not(None))
        .order_by(HRTemplate.version.desc())
        .limit(1)
    )
    template = result.scalar_one_or_none()
    if not template:
        return None
    return TemplateDetailResponse.model_validate(template)


# ---------------------------------------------------------------------------
# Deterministic template resolution (for the HR chat tool)
# ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class TemplateCandidate:
    """Lightweight, DB-free view of a render-eligible template."""

    id: int
    category: str
    name: str


@dataclass(frozen=True, slots=True)
class ResolveResult:
    """Outcome of :func:`resolve_template`.

    ``match`` is the single unambiguous template, or ``None`` when the query is
    ambiguous / matched nothing. ``candidates`` is every render-eligible template
    (used to offer a list; empty means the library has nothing renderable).
    """

    match: TemplateDetailResponse | None
    candidates: list[TemplateCandidate]


def _normalize(text: str) -> str:
    """Lowercase, trim, collapse runs of whitespace/underscores to one space."""
    return re.sub(r"[\s_]+", " ", text.strip().lower())


def _query_matches(normalized_query: str, candidate: TemplateCandidate) -> bool:
    """True if the query is contained in (or contains) the category or name."""
    category = _normalize(candidate.category)
    name = _normalize(candidate.name)
    return (
        normalized_query in category
        or category in normalized_query
        or normalized_query in name
        or name in normalized_query
    )


def match_template(
    candidates: list[TemplateCandidate],
    query: str,
) -> TemplateCandidate | None:
    """Resolve ``query`` to exactly one candidate, deterministically.

    Priority, highest first: exact normalized category slug, then exact
    normalized name, then normalized containment. Returns a candidate only when
    the highest level that matches anything matches EXACTLY one; an empty query,
    no match, or an ambiguous (>1) level returns ``None``. Pure function: no DB,
    no randomness, no LLM.
    """
    normalized_query = _normalize(query)
    if not normalized_query:
        return None

    for exact in (
        [c for c in candidates if _normalize(c.category) == normalized_query],
        [c for c in candidates if _normalize(c.name) == normalized_query],
    ):
        if len(exact) == 1:
            return exact[0]
        if exact:  # more than one at this level -> ambiguous
            return None

    contains = [c for c in candidates if _query_matches(normalized_query, c)]
    return contains[0] if len(contains) == 1 else None


async def resolve_template(db: AsyncSession, query: str) -> ResolveResult:
    """Resolve a free-text query to one render-eligible template, or a list.

    Loads the active, render-eligible templates (``is_active`` + non-empty
    ``template_fields``), runs the deterministic matcher, and returns either the
    single matched template (full detail) or the candidate list.
    """
    result = await db.execute(
        select(HRTemplate)
        .where(HRTemplate.is_active.is_(True))
        .where(HRTemplate.template_fields.is_not(None))
        .order_by(HRTemplate.category)
    )
    rows = [row for row in result.scalars().all() if row.template_fields]
    candidates = [
        TemplateCandidate(id=row.id, category=row.category, name=row.name)
        for row in rows
    ]

    matched = match_template(candidates, query)
    if matched is None:
        return ResolveResult(match=None, candidates=candidates)

    matched_row = next(row for row in rows if row.id == matched.id)
    return ResolveResult(
        match=TemplateDetailResponse.model_validate(matched_row),
        candidates=candidates,
    )


async def list_template_history(
    db: AsyncSession,
    category: str,
) -> list[TemplateResponse]:
    """Return all template versions for a category, newest first."""
    result = await db.execute(
        select(HRTemplate)
        .where(HRTemplate.category == category)
        .order_by(HRTemplate.version.desc())
    )
    items = list(result.scalars().all())
    return [TemplateResponse.model_validate(t) for t in items]


async def rollback_template(
    db: AsyncSession,
    template_id: int,
) -> HRTemplate:
    """Rollback category to the specified template version.

    Deactivates the current active template and reactivates the target.
    """
    target_result = await db.execute(
        select(HRTemplate).where(HRTemplate.id == template_id)
    )
    target = target_result.scalar_one_or_none()
    if target is None:
        raise HRFileError("Template not found")

    current_result = await db.execute(
        select(HRTemplate)
        .where(HRTemplate.category == target.category)
        .where(HRTemplate.is_active.is_(True))
        .with_for_update()
    )
    current = current_result.scalar_one_or_none()

    if current is not None:
        current.is_active = False
        current.superseded_by_id = target.id
        await db.flush()

    target.is_active = True
    target.superseded_by_id = None
    await db.flush()
    await db.refresh(target)

    return target


async def delete_template(
    db: AsyncSession,
    template_id: int,
) -> bool:
    """Delete template and its file."""
    result = await db.execute(select(HRTemplate).where(HRTemplate.id == template_id))
    template = result.scalar_one_or_none()
    if not template:
        return False

    try:
        Path(template.file_path).unlink(missing_ok=True)
    except OSError:
        logger.warning("Could not delete file: %s", template.file_path)

    await db.delete(template)
    return True
