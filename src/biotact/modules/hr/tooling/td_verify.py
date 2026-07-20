"""Verify a hand-rebuilt TD template against the original and the target rules.

Four checks (see spec 2026-07-20-hr-pr-b): text 1:1, placeholders, clean
numbering, row-per-clause structure. Pure functions over ``python-docx`` — no
network, no DB. Run: ``python -m biotact.modules.hr.tooling.td_verify ORIG CAND``.
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from docx import Document
from docx.oxml.ns import qn

from biotact.modules.hr.tooling.docx_text import (
    clause_bodies,
    extract_columns,
    placeholders,
    starts_with_clause_number,
    strip_leading_clause_number,
    walk_paragraphs,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

_PREVIEW = 80


@dataclass
class CheckResult:
    """Outcome of one verification check."""

    name: str
    ok: bool
    problems: list[str] = field(default_factory=list)


def _column_diff(lang: str, expected: list[str], actual: list[str]) -> list[str]:
    """Ordered-sequence diff of one language column (first divergence + count)."""
    problems: list[str] = []
    if len(expected) != len(actual):
        problems.append(f"{lang}: {len(actual)} clauses, expected {len(expected)}")
    for index, (want, got) in enumerate(zip(expected, actual, strict=False)):
        if want != got:
            problems.append(
                f"{lang}: clause {index} differs — "
                f"expected {want[:_PREVIEW]!r}, got {got[:_PREVIEW]!r}"
            )
            break  # a shift cascades; the first divergence is the signal
    return problems


def check_text_1to1(original: DocxDocument, candidate: DocxDocument) -> CheckResult:
    """Each language column's clause sequence is preserved 1:1 AND in order.

    Compared as ORDERED per-column sequences (not multisets), so a clause moved
    into another section is caught, not just lost/added/altered text. Numbers
    are stripped from the ORIGINAL only (the target has none in its text); a
    candidate body still starting with a literal number is flagged as an author
    typing a digit — the exact defect PR-B removes.
    """
    problems: list[str] = []
    original_columns = extract_columns(original)
    candidate_columns = extract_columns(candidate)

    for lang, original_column, candidate_column in (
        ("col0", original_columns[0], candidate_columns[0]),
        ("col1", original_columns[1], candidate_columns[1]),
    ):
        for body in candidate_column:
            if starts_with_clause_number(body):
                problems.append(
                    f"{lang}: candidate typed a literal number: {body[:_PREVIEW]!r}"
                )
        expected = [strip_leading_clause_number(body) for body in original_column]
        problems.extend(_column_diff(lang, expected, candidate_column))

    return CheckResult("text_1to1", not problems, problems)


def check_placeholders(original: DocxDocument, candidate: DocxDocument) -> CheckResult:
    """Candidate's ``{{ }}`` placeholder set equals the original's."""
    problems: list[str] = []
    original_names = placeholders(original)
    candidate_names = placeholders(candidate)
    missing = original_names - candidate_names
    unexpected = candidate_names - original_names
    if missing:
        problems.append(f"missing placeholders: {sorted(missing)}")
    if unexpected:
        problems.append(f"unexpected placeholders: {sorted(unexpected)}")
    return CheckResult("placeholders", not problems, problems)


def _has_literal_digit(lvl_text: str | None) -> bool:
    """True if ``lvlText`` hardcodes a digit (a digit not part of a ``%N`` token)."""
    if lvl_text is None:
        return False
    return any(char.isdigit() for char in re.sub(r"%\d", "", lvl_text))


def _num_id_from_ppr(ppr: object) -> str | None:
    """Return the numId in a ``<w:pPr>`` (paragraph or style), or ``None``."""
    if ppr is None:
        return None
    num_pr = ppr.find(qn("w:numPr"))  # type: ignore[attr-defined]
    if num_pr is None:
        return None
    num_id = num_pr.find(qn("w:numId"))
    return num_id.get(qn("w:val")) if num_id is not None else None


def _paragraph_num_id(paragraph: object) -> str | None:
    """numId from the paragraph directly, else inherited from its style chain."""
    direct = _num_id_from_ppr(paragraph._p.find(qn("w:pPr")))  # type: ignore[attr-defined]
    if direct is not None:
        return direct
    style = paragraph.style  # type: ignore[attr-defined]
    while style is not None:
        inherited = _num_id_from_ppr(style.element.find(qn("w:pPr")))
        if inherited is not None:
            return inherited
        style = style.base_style
    return None


def _used_abstract_num_ids(
    candidate: DocxDocument, num_to_abstract: dict[str, str | None]
) -> set[str]:
    """AbstractNumIds actually referenced by numbered paragraphs (direct or via style)."""
    used: set[str] = set()
    for paragraph in walk_paragraphs(candidate):
        num_id = _paragraph_num_id(paragraph)
        if num_id is None:
            continue
        abstract = num_to_abstract.get(num_id)
        if abstract is not None:
            used.add(abstract)
    return used


def check_numbering(candidate: DocxDocument) -> CheckResult:
    """Numbered clauses use ONE clean multilevel list.

    Passes when: all numbered paragraphs resolve to a single ``abstractNum``
    (many ``numId`` pointing at it is fine — Word does that legitimately);
    no ``lvlOverride``/``startOverride``; lvl0 is ``decimal``; every level
    ``start=1``; no literal digits in any ``lvlText``.
    """
    try:
        numbering = candidate.part.numbering_part.element
    except (NotImplementedError, AttributeError):
        return CheckResult(
            "numbering", False, ["no numbering part — document has no numbered list"]
        )

    problems: list[str] = []
    num_to_abstract: dict[str, str | None] = {}
    for num in numbering.findall(qn("w:num")):
        abstract = num.find(qn("w:abstractNumId"))
        num_to_abstract[num.get(qn("w:numId"))] = (
            abstract.get(qn("w:val")) if abstract is not None else None
        )
        if num.find(qn("w:lvlOverride")) is not None:
            problems.append(
                f"num {num.get(qn('w:numId'))} has a lvlOverride/startOverride"
            )

    used = _used_abstract_num_ids(candidate, num_to_abstract)
    if not used:
        return CheckResult("numbering", False, ["no numbered clauses found"])
    if len(used) > 1:
        problems.append(
            f"clauses use {len(used)} numbering definitions, expected 1: {sorted(used)}"
        )

    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) not in used:
            continue
        aid = abstract.get(qn("w:abstractNumId"))
        for lvl in abstract.findall(qn("w:lvl")):
            ilvl = lvl.get(qn("w:ilvl"))
            fmt = lvl.find(qn("w:numFmt"))
            start = lvl.find(qn("w:start"))
            lvl_text = lvl.find(qn("w:lvlText"))
            if ilvl == "0" and fmt is not None and fmt.get(qn("w:val")) != "decimal":
                problems.append(
                    f"abstractNum {aid} lvl0 numFmt={fmt.get(qn('w:val'))}, expected decimal"
                )
            if start is not None and start.get(qn("w:val")) != "1":
                problems.append(
                    f"abstractNum {aid} lvl{ilvl} start={start.get(qn('w:val'))}, expected 1"
                )
            text_val = lvl_text.get(qn("w:val")) if lvl_text is not None else None
            if _has_literal_digit(text_val):
                problems.append(
                    f"abstractNum {aid} lvl{ilvl} lvlText={text_val!r} hardcodes a digit"
                )

    return CheckResult("numbering", not problems, problems)


def check_structure(
    candidate: DocxDocument, empty_cell_rows: frozenset[int] = frozenset()
) -> CheckResult:
    """One 2-column table; each row has UZ and RU text (row = one clause).

    A cell may be empty only for a row index in ``empty_cell_rows`` (the
    ratified list of clauses that legitimately exist in one language only).
    """
    problems: list[str] = []
    tables = candidate.tables
    if len(tables) != 1:
        return CheckResult(
            "structure", False, [f"expected exactly 1 table, found {len(tables)}"]
        )
    table = tables[0]
    if len(table.columns) != 2:
        problems.append(f"expected 2 columns, found {len(table.columns)}")
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells[:2]):
            if cell.tables:
                problems.append(
                    f"row {row_index} col {col_index} contains a nested table "
                    "(cells must be flat, one clause per row)"
                )
            if not clause_bodies(cell) and row_index not in empty_cell_rows:
                problems.append(
                    f"row {row_index} col {col_index} is empty (not in exception list)"
                )
    return CheckResult("structure", not problems, problems)


def verify(
    original_path: str,
    candidate_path: str,
    *,
    empty_cell_rows: frozenset[int] = frozenset(),
) -> list[CheckResult]:
    """Run all four checks and return their results."""
    original = Document(original_path)
    candidate = Document(candidate_path)
    return [
        check_text_1to1(original, candidate),
        check_placeholders(original, candidate),
        check_numbering(candidate),
        check_structure(candidate, empty_cell_rows),
    ]


def _main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("original", help="path to the current (original) template")
    parser.add_argument("candidate", help="path to the hand-rebuilt candidate")
    parser.add_argument(
        "--empty-rows",
        default="",
        help="comma-separated row indices allowed to have an empty cell",
    )
    args = parser.parse_args(argv)
    empty_rows = frozenset(
        int(part) for part in args.empty_rows.split(",") if part.strip()
    )
    results = verify(args.original, args.candidate, empty_cell_rows=empty_rows)
    for result in results:
        status = "OK  " if result.ok else "FAIL"
        print(f"[{status}] {result.name}")
        for problem in result.problems:
            print(f"        - {problem}")
    return 0 if all(result.ok for result in results) else 1


if __name__ == "__main__":
    raise SystemExit(_main())
