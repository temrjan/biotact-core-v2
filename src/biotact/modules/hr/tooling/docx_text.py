"""Shared DOCX text helpers for the TD template tooling.

Pure functions over ``python-docx`` objects: recursive paragraph walk (the
template nests 6-7 tables inside each language cell, which ``cell.paragraphs``
does not reach), text normalization, placeholder scan, and clause-number
handling.
"""

from __future__ import annotations

import re
import unicodedata
from typing import TYPE_CHECKING

from docx.table import Table
from docx.text.paragraph import Paragraph

if TYPE_CHECKING:
    from collections.abc import Iterator

PLACEHOLDER_RE = re.compile(r"\{\{\s*(\w+)\s*\}\}")

# A leading clause number: 1-3 dotted groups of 1-2 digits, ending in a dot +
# whitespace ("2. ", "2.1. ", "2.1.3. "). Deliberately NOT a date: "05.01.2026 "
# has a 4-digit tail and no trailing "dot + space", so it never matches.
CLAUSE_NUMBER_RE = re.compile(r"^\d{1,2}(?:\.\d{1,2}){0,2}\.\s+")

_WHITESPACE_RE = re.compile(r"\s+")
_NBSP = " "


def walk_paragraphs(container: object) -> Iterator[Paragraph]:
    """Yield every paragraph in document order, descending into nested tables.

    ``container`` is a ``Document`` or a table ``_Cell`` — both expose
    ``iter_inner_content()``.
    """
    for item in container.iter_inner_content():  # type: ignore[attr-defined]
        if isinstance(item, Paragraph):
            yield item
        elif isinstance(item, Table):
            for row in item.rows:
                for cell in row.cells:
                    yield from walk_paragraphs(cell)


def normalize(text: str) -> str:
    """Normalize for text-equality: NFC, NBSP->space, collapse whitespace, strip.

    Case, quotes and dashes are left untouched on purpose — in legal text their
    substitution is a real difference and must surface as a mismatch.
    """
    text = unicodedata.normalize("NFC", text)
    text = text.replace(_NBSP, " ")
    return _WHITESPACE_RE.sub(" ", text).strip()


def placeholders(container: object) -> set[str]:
    """Return the set of ``{{ NAME }}`` placeholder names in the document."""
    names: set[str] = set()
    for paragraph in walk_paragraphs(container):
        names.update(PLACEHOLDER_RE.findall(paragraph.text))
    return names


def starts_with_clause_number(text: str) -> bool:
    """True if ``text`` begins with a literal clause number (e.g. ``2.1. ``)."""
    return CLAUSE_NUMBER_RE.match(text) is not None


def strip_leading_clause_number(text: str) -> str:
    """Drop a single leading clause number if present; otherwise return as-is."""
    return CLAUSE_NUMBER_RE.sub("", text, count=1)


def clause_bodies(container: object) -> list[str]:
    """Normalized, non-empty paragraph texts in document order."""
    bodies = [normalize(p.text) for p in walk_paragraphs(container)]
    return [body for body in bodies if body]


def extract_columns(doc: object) -> tuple[list[str], list[str]]:
    """Ordered (col0, col1) clause bodies from the document's first table.

    In the TD template col0 is UZ and col1 is RU; the order within a column is
    preserved so callers can compare sequences (clause order), not just sets.
    """
    table = doc.tables[0]  # type: ignore[attr-defined]
    left: list[str] = []
    right: list[str] = []
    for row in table.rows:
        cells = row.cells
        if len(cells) >= 1:
            left.extend(clause_bodies(cells[0]))
        if len(cells) >= 2:
            right.extend(clause_bodies(cells[1]))
    return left, right
