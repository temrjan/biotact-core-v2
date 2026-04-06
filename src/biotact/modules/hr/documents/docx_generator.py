"""DOCX generator — convert text to downloadable DOCX file."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt


def text_to_docx(text: str, title: str | None = None) -> io.BytesIO:
    """Convert plain/markdown-like text to DOCX.

    Handles:
    - Lines starting with # as headings
    - Bold (**text**) markers
    - Regular paragraphs
    - Numbered and bullet lists

    Returns: BytesIO buffer with DOCX content.
    """
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    if title:
        doc.add_heading(title, level=0)

    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue

        # Headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        # Bullet lists
        elif stripped.startswith("- ") or stripped.startswith("• "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        # Numbered lists
        elif re.match(r"^\d+\.\s", stripped):
            text_content = re.sub(r"^\d+\.\s+", "", stripped)
            doc.add_paragraph(text_content, style="List Number")
        # Regular paragraph
        else:
            p = doc.add_paragraph()
            # Handle bold markers
            parts = re.split(r"(\*\*[^*]+\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
