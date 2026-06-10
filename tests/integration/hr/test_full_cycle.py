"""Integration tests for HR full cycle: upload → scan → render → download (PR-11)."""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

import pytest
import pytest_asyncio

from biotact.core.config import Settings, get_settings
from biotact.main import app
from biotact.modules.hr.library.models import HRDocument

if TYPE_CHECKING:
    from collections.abc import Generator

    from httpx import AsyncClient
    from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.integration

MINIMAL_DOCX = b"PK\x03\x04" + b"\x00" * 200


def _hr_settings() -> Settings:
    return Settings(
        app_env="development",
        debug=False,
        secret_key="test-secret-key-for-testing-only",
        postgres_host="localhost",
        postgres_port=5432,
        postgres_user="biotact",
        postgres_password="biotact_test",
        postgres_db="biotact_test",
        openai_api_key="sk-test-key",
        hr_allowed_emails="test@biotact.uz",
    )


@pytest.fixture
def hr_allow() -> Generator[None, None, None]:
    app.dependency_overrides[get_settings] = _hr_settings
    yield
    app.dependency_overrides.pop(get_settings, None)


@pytest_asyncio.fixture
async def hr_client(
    authenticated_client: AsyncClient,
    hr_allow: None,
) -> AsyncClient:
    return authenticated_client


def _make_docx_with_placeholders(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Employee: {{ FIO }}")
    doc.add_paragraph("Position: {{ POSITION }}")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Salary: {{ SALARY }}"
    doc.save(str(path))


class TestFullCycle:
    """Upload → scan → render → download."""

    @pytest.mark.asyncio
    async def test_upload_extracts_fields(
        self, hr_client: AsyncClient, tmp_path: Path
    ) -> None:
        docx = tmp_path / "template.docx"
        _make_docx_with_placeholders(docx)

        with docx.open("rb") as f:
            files = {
                "file": (
                    "template.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
            r = await hr_client.post(
                "/api/v1/hr/library?category=td_cycle", files=files
            )

        assert r.status_code == 201
        data = r.json()
        assert data["category"] == "td_cycle"
        assert set(data.get("template_fields") or []) == {"FIO", "POSITION", "SALARY"}

    @pytest.mark.asyncio
    async def test_render_returns_docx(
        self, hr_client: AsyncClient, tmp_path: Path
    ) -> None:
        docx = tmp_path / "template.docx"
        _make_docx_with_placeholders(docx)

        with docx.open("rb") as f:
            files = {
                "file": (
                    "template.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            }
            r = await hr_client.post(
                "/api/v1/hr/library?category=td_render", files=files
            )
        assert r.status_code == 201
        tpl_id = r.json()["id"]

        payload = {
            "template_id": tpl_id,
            "data": {"FIO": "Иванов Иван", "POSITION": "Менеджер", "SALARY": "5000000"},
            "filename": "contract.docx",
        }
        render = await hr_client.post("/api/v1/hr/documents/render", json=payload)
        assert render.status_code == 200
        assert (
            render.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert b"PK" in render.content

    @pytest.mark.asyncio
    async def test_download_existing_document(
        self, hr_client: AsyncClient, test_session: AsyncSession
    ) -> None:
        """Create HRDocument row + file, then download via API."""
        from biotact.modules.hr.library.models import HRTemplate

        tpl = HRTemplate(
            name="dl.docx",
            category="td_dl",
            file_path="data/hr_templates/dl.docx",
            file_type="docx",
            file_size=100,
            version=1,
            is_active=True,
            uploaded_by=1,
        )
        test_session.add(tpl)
        await test_session.flush()

        render_dir = Path("data/hr_rendered")
        render_dir.mkdir(parents=True, exist_ok=True)
        file_path = render_dir / "dltest.docx"
        file_path.write_bytes(MINIMAL_DOCX)

        doc = HRDocument(
            file_id="dltest",
            template_id=tpl.id,
            template_name="dl.docx",
            employee_name="Test",
            file_path=str(file_path),
            file_size=len(MINIMAL_DOCX),
            created_by=1,
        )
        test_session.add(doc)
        await test_session.commit()

        r = await hr_client.get("/api/v1/hr/documents/download/dltest")
        assert r.status_code == 200
        assert (
            r.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
