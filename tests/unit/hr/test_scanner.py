"""Unit tests for DOCX template field scanner."""

from __future__ import annotations

from typing import TYPE_CHECKING

from biotact.modules.hr.library.scanner import scan_template_fields

if TYPE_CHECKING:
    from pathlib import Path


def _make_docx_with_text(path: Path, text: str) -> None:
    """Create a minimal DOCX containing the given text."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def _make_docx_with_table(path: Path, cell_text: str) -> None:
    """Create a DOCX with a table cell containing text."""
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = cell_text
    doc.save(str(path))


class TestScanTemplateFields:
    """scan_template_fields extracts {{ PLACEHOLDER }} names."""

    def test_single_placeholder(self, tmp_path: Path) -> None:
        docx = tmp_path / "single.docx"
        _make_docx_with_text(docx, "Hello {{ FIO }}")
        fields = scan_template_fields(str(docx))
        assert fields == ["FIO"]

    def test_multiple_placeholders(self, tmp_path: Path) -> None:
        docx = tmp_path / "multi.docx"
        _make_docx_with_text(docx, "{{ FIO }} works as {{ POSITION }}")
        fields = scan_template_fields(str(docx))
        assert fields == ["FIO", "POSITION"]

    def test_placeholder_in_table(self, tmp_path: Path) -> None:
        docx = tmp_path / "table.docx"
        _make_docx_with_table(docx, "Salary: {{ SALARY }}")
        fields = scan_template_fields(str(docx))
        assert fields == ["SALARY"]

    def test_whitespace_inside_braces(self, tmp_path: Path) -> None:
        docx = tmp_path / "spaces.docx"
        _make_docx_with_text(docx, "{{  FIO  }}")
        fields = scan_template_fields(str(docx))
        assert fields == ["FIO"]

    def test_no_placeholders(self, tmp_path: Path) -> None:
        docx = tmp_path / "empty.docx"
        _make_docx_with_text(docx, "Just plain text")
        fields = scan_template_fields(str(docx))
        assert fields == []

    def test_nonexistent_file_returns_empty(self) -> None:
        fields = scan_template_fields("/tmp/does_not_exist.docx")
        assert fields == []
