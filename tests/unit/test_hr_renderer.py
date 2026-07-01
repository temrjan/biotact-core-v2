"""Unit tests for HR DOCX table-width normalization.

Locks the fix for "поплывшие ячейки": after ``render_template`` every table is
clamped to the page width with a fixed layout and integer widths, so MS Word
never draws a table wider than the page. Fixtures are built in memory (no binary
.docx in the repo). See ``.claude/specs/2026-06-30-hr-docx-table-width-normalization.md``.
"""

from __future__ import annotations

import re

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from biotact.modules.hr.documents.renderer import (
    _normalize_one_table,
    _normalize_table_widths,
    _scale_down,
    _text_width_twips,
    render_template,
)

pytestmark = pytest.mark.unit

W = qn("w:w")
TYPE = qn("w:type")
VAL = qn("w:val")

# Letter page 12240 twips, 1800 margins each side -> 8640 usable text width.
PAGE = 12240
MARGIN = 1800
TEXT_WIDTH = PAGE - 2 * MARGIN  # 8640


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------


def _blank_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Twips(PAGE)
    section.left_margin = Twips(MARGIN)
    section.right_margin = Twips(MARGIN)
    return doc


def _raw_table(
    grid_widths: list[str | None],
    rows: list[list[int]] | None = None,
    *,
    tbl_ind: int | None = None,
    with_tbl_pr: bool = True,
    with_tc_pr: bool = True,
) -> OxmlElement:
    """Build a ``<w:tbl>`` element with the given grid and row cell spans.

    Args:
        grid_widths: Raw ``w:w`` strings for each ``<w:gridCol>``; ``None`` omits
            the ``w:w`` attribute (a malformed column).
        rows: Per row, a list of ``gridSpan`` values (one per cell). Defaults to
            a single row with one span-1 cell per column.
        tbl_ind: Optional table indent (``<w:tblInd>``) in twips.
        with_tbl_pr: Emit a ``<w:tblPr>`` (set False for the malformed case).
        with_tc_pr: Emit a ``<w:tcPr>`` per cell (set False for the malformed case).
    """
    tbl = OxmlElement("w:tbl")

    if with_tbl_pr:
        tbl_pr = OxmlElement("w:tblPr")
        if tbl_ind is not None:
            ind = OxmlElement("w:tblInd")
            ind.set(W, str(tbl_ind))
            ind.set(TYPE, "dxa")
            tbl_pr.append(ind)
        tbl.append(tbl_pr)

    grid = OxmlElement("w:tblGrid")
    for width in grid_widths:
        col = OxmlElement("w:gridCol")
        if width is not None:
            col.set(W, width)
        grid.append(col)
    tbl.append(grid)

    if rows is None:
        rows = [[1] * len(grid_widths)]
    for spans in rows:
        tr = OxmlElement("w:tr")
        for span in spans:
            tc = OxmlElement("w:tc")
            if with_tc_pr:
                tc_pr = OxmlElement("w:tcPr")
                if span > 1:
                    grid_span = OxmlElement("w:gridSpan")
                    grid_span.set(VAL, str(span))
                    tc_pr.append(grid_span)
                tc.append(tc_pr)
            tc.append(OxmlElement("w:p"))
            tr.append(tc)
        tbl.append(tr)

    return tbl


def _attach(doc: Document, tbl: OxmlElement) -> OxmlElement:
    """Insert a raw table before the body's section properties."""
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    sect_pr.addprevious(tbl)
    return tbl


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------


def _grid(tbl: OxmlElement) -> list[int]:
    grid = tbl.find(qn("w:tblGrid"))
    return [int(col.get(W)) for col in grid.findall(qn("w:gridCol"))]


def _raw_grid(tbl: OxmlElement) -> list[str | None]:
    grid = tbl.find(qn("w:tblGrid"))
    return [col.get(W) for col in grid.findall(qn("w:gridCol"))]


def _layout(tbl: OxmlElement) -> str | None:
    tbl_pr = tbl.find(qn("w:tblPr"))
    if tbl_pr is None:
        return None
    layout = tbl_pr.find(qn("w:tblLayout"))
    return layout.get(TYPE) if layout is not None else None


def _tbl_w(tbl: OxmlElement) -> tuple[str | None, str | None]:
    tbl_pr = tbl.find(qn("w:tblPr"))
    tbl_w = tbl_pr.find(qn("w:tblW")) if tbl_pr is not None else None
    if tbl_w is None:
        return (None, None)
    return (tbl_w.get(W), tbl_w.get(TYPE))


def _cell_widths(tbl: OxmlElement, row: int = 0) -> list[int]:
    widths: list[int] = []
    for tc in tbl.findall(qn("w:tr"))[row].findall(qn("w:tc")):
        tc_w = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
        widths.append(int(tc_w.get(W)))
    return widths


def _has_fractional(tbl: OxmlElement) -> bool:
    return re.search(r'w:w="\d+\.\d+"', tbl.xml) is not None


# ===========================================================================
# 1. autofit + oversize -> clamp to page, fixed layout, proportions kept
# ===========================================================================


def test_autofit_oversize_clamped_to_page() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["7000", "7000"]))

    _normalize_table_widths(doc)

    assert _layout(tbl) == "fixed"
    assert _tbl_w(tbl) == (str(TEXT_WIDTH), "dxa")
    assert _grid(tbl) == [4320, 4320]
    assert sum(_grid(tbl)) == TEXT_WIDTH


def test_oversize_preserves_proportions() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["7000", "3000"]))  # 70/30 split, sum 10000

    _normalize_table_widths(doc)

    # factor = 8640/10000 = 0.864 -> 6048 / 2592, sum exactly 8640
    assert _grid(tbl) == [6048, 2592]
    assert sum(_grid(tbl)) == TEXT_WIDTH


# ===========================================================================
# 2. fractional widths -> rounded to integers
# ===========================================================================


def test_fractional_widths_rounded_when_narrow() -> None:
    # 4785 + 125 = 4910 < 8640 -> only-shrink keeps widths, just rounds them.
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["4785.0", "124.99999999999909"]))

    _normalize_table_widths(doc)

    assert _grid(tbl) == [4785, 125]
    assert not _has_fractional(tbl)


def test_fractional_widths_rounded_when_oversize() -> None:
    # ТД case: fractional widths whose sum is ~3x the page.
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["8639.9999", "8639.9999", "8639.9999"]))

    _normalize_table_widths(doc)

    assert not _has_fractional(tbl)
    assert sum(_grid(tbl)) == TEXT_WIDTH
    assert _layout(tbl) == "fixed"


# ===========================================================================
# 3. gridSpan -> cell width = sum of covered columns
# ===========================================================================


def test_gridspan_cell_width_sums_columns() -> None:
    # Grid sums to 8640 exactly (narrow path, no scaling), one span-2 cell.
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["3000", "3000", "2640"], rows=[[2, 1]]))

    _normalize_table_widths(doc)

    assert _cell_widths(tbl) == [6000, 2640]  # cols 0+1, then col 2


# ===========================================================================
# 4. narrow table -> widths unchanged (only rounding), still fixed
# ===========================================================================


def test_narrow_table_not_expanded() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["2000", "2000"]))

    _normalize_table_widths(doc)

    assert _grid(tbl) == [2000, 2000]
    assert _layout(tbl) == "fixed"
    assert _tbl_w(tbl) == ("4000", "dxa")


# ===========================================================================
# 5. tblInd affects the clamp target
# ===========================================================================


def test_tblind_reduces_max_width() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["8000", "8000"], tbl_ind=1000))

    _normalize_table_widths(doc)

    assert sum(_grid(tbl)) == TEXT_WIDTH - 1000  # 7640


def test_large_positive_tblind_falls_back_to_text_width() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["8000", "8000"], tbl_ind=9000))

    _normalize_table_widths(doc)

    # max_width would be negative -> fall back to full text width.
    assert sum(_grid(tbl)) == TEXT_WIDTH


# ===========================================================================
# 6. guards — no exception, MINOR-1 (no forced fixed in fallbacks)
# ===========================================================================


def test_document_without_tables_is_noop() -> None:
    doc = _blank_doc()
    doc.add_paragraph("just text")
    before = doc.element.body.xml

    _normalize_table_widths(doc)

    assert doc.element.body.xml == before
    assert doc.element.body.find(".//" + qn("w:tbl")) is None


def test_zero_sum_grid_rounds_only_no_fixed() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["0", "0"]))

    _normalize_table_widths(doc)

    assert _layout(tbl) is None  # fallback: layout mode left untouched
    assert _tbl_w(tbl) == (None, None)


def test_text_width_none_rounds_only_no_fixed() -> None:
    # No clamp target -> round widths, leave layout mode and scaling untouched.
    tbl = _raw_table(["8639.9999", "8639.9999"])

    _normalize_one_table(tbl, None)

    assert not _has_fractional(tbl)
    assert _grid(tbl) == [8640, 8640]  # rounded, not scaled
    assert _layout(tbl) is None


# ===========================================================================
# 7. nested tables normalized too
# ===========================================================================


def test_nested_table_is_normalized() -> None:
    doc = _blank_doc()
    outer = _attach(doc, _raw_table(["7000", "7000"]))
    inner = _raw_table(["9000", "9000"])
    # nest the inner table inside the first cell of the outer table
    outer.findall(qn("w:tr"))[0].findall(qn("w:tc"))[0].append(inner)

    _normalize_table_widths(doc)

    assert _layout(outer) == "fixed"
    assert _layout(inner) == "fixed"
    assert sum(_grid(inner)) == TEXT_WIDTH


# ===========================================================================
# 8. idempotent
# ===========================================================================


def test_normalization_is_idempotent() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["7000", "3000"], rows=[[2]]))

    _normalize_table_widths(doc)
    first = tbl.xml
    _normalize_table_widths(doc)

    assert tbl.xml == first


# ===========================================================================
# 9. _text_width_twips helper
# ===========================================================================


def test_text_width_twips_single_section() -> None:
    doc = _blank_doc()
    assert _text_width_twips(doc) == TEXT_WIDTH


def test_text_width_twips_returns_narrowest_section() -> None:
    doc = _blank_doc()
    new = doc.add_section()
    new.page_width = Twips(20000)
    new.left_margin = Twips(1000)
    new.right_margin = Twips(1000)  # this section text width = 18000
    # narrowest is the original 8640
    assert _text_width_twips(doc) == TEXT_WIDTH


# ===========================================================================
# 10. end-to-end render_template (locks the doc.docx path vs get_docx regression)
# ===========================================================================


def test_render_template_replaces_and_normalizes(tmp_path) -> None:
    doc = _blank_doc()
    doc.add_paragraph("Employee: {{ FIO }}")
    table = doc.add_table(rows=1, cols=2)
    for col in table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")):
        col.set(W, "24999.999999")  # oversize + fractional
    template_path = tmp_path / "contract.docx"
    doc.save(str(template_path))

    buffer = render_template(str(template_path), {"FIO": "Ivan Petrov"})
    out = Document(buffer)

    # render actually happened (guards against the get_docx() reload regression)
    assert "Ivan Petrov" in out.paragraphs[0].text
    # normalization happened
    assert not re.search(r'w:w="\d+\.\d+"', out.element.body.xml)
    tbl = out.element.body.find(".//" + qn("w:tbl"))
    assert _layout(tbl) == "fixed"
    assert sum(_grid(tbl)) <= TEXT_WIDTH


# ===========================================================================
# 11. _scale_down remainder distribution (exact sum under integer rounding)
# ===========================================================================


@pytest.mark.parametrize(
    ("widths", "target"),
    [
        ([1000, 1000, 1000], 2000),  # rounds up -> must decrement (diff < 0)
        ([1000, 1000, 1000], 2002),  # rounds down -> must increment (diff > 0)
        ([7000, 3000], 8640),  # clean division (diff == 0)
        ([3333, 3333, 3334], 8639),  # odd target, mixed rounding remainder
        ([5, 5, 5, 5, 5], 3),  # target below the 1-twip floor
    ],
)
def test_scale_down_sums_exactly_to_target(widths: list[int], target: int) -> None:
    scaled = _scale_down(widths, target)

    assert len(scaled) == len(widths)
    assert all(width >= 1 for width in scaled)
    # Exactly the target, unless the target is below the 1-twip-per-column floor,
    # in which case the columns settle at the floor (guard prevents an infinite loop).
    assert sum(scaled) == max(target, len(widths))


def test_scale_down_preserves_proportions() -> None:
    widths = [3333, 3333, 3334]
    total = sum(widths)

    scaled = _scale_down(widths, 8639)

    for original, result in zip(widths, scaled, strict=True):
        assert abs(result - round(original * 8639 / total)) <= 1


# ===========================================================================
# 12. gridSpan exceeding columns -> clamped, no index error
# ===========================================================================


def test_gridspan_exceeding_columns_is_clamped() -> None:
    doc = _blank_doc()
    # one cell spanning more columns than exist -> clamped to the whole grid
    over = _attach(doc, _raw_table(["3000", "3000", "2640"], rows=[[5]]))
    # a cell that begins past the last column -> zero width, no index error
    past = _attach(doc, _raw_table(["3000", "3000", "2640"], rows=[[3, 2]]))

    _normalize_table_widths(doc)

    assert _cell_widths(over) == [TEXT_WIDTH]
    assert _cell_widths(past) == [TEXT_WIDTH, 0]


# ===========================================================================
# 13. multi-row cell widths — per-row cursor reset
# ===========================================================================


def test_multi_row_cell_widths_reset_per_row() -> None:
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["3000", "3000", "2640"], rows=[[2, 1], [1, 1, 1]]))

    _normalize_table_widths(doc)

    assert _cell_widths(tbl, row=0) == [6000, 2640]
    assert _cell_widths(tbl, row=1) == [3000, 3000, 2640]


# ===========================================================================
# 14. malformed elements — fail-soft create/None branches, no crash
# ===========================================================================


def test_malformed_table_elements_are_handled() -> None:
    doc = _blank_doc()
    # oversize table with no <w:tblPr>, a column missing w:w, and cells with no <w:tcPr>
    tbl = _attach(
        doc,
        _raw_table(["7000", None, "7000"], with_tbl_pr=False, with_tc_pr=False),
    )

    _normalize_table_widths(doc)  # exercises the create / None fail-soft branches

    assert _layout(tbl) == "fixed"  # tblPr created, fixed layout applied
    assert sum(_grid(tbl)) == TEXT_WIDTH  # missing-width column treated as 0


def test_non_numeric_width_treated_as_zero() -> None:
    # OOXML w:w may be a universal measure ("1in"); it must not crash the table.
    doc = _blank_doc()
    tbl = _attach(doc, _raw_table(["1in", "5000", "5000"]))

    _normalize_table_widths(doc)

    assert not _has_fractional(tbl)
    assert _layout(tbl) == "fixed"
    assert sum(_grid(tbl)) == TEXT_WIDTH


def test_section_without_margins_yields_no_text_width() -> None:
    doc = _blank_doc()
    # strip <w:pgMar> so the section reports None margins
    sect_pr = doc.element.body.find(qn("w:sectPr"))
    sect_pr.remove(sect_pr.find(qn("w:pgMar")))

    assert _text_width_twips(doc) is None


def test_normalization_is_fail_soft_on_error() -> None:
    class Boom:
        @property
        def sections(self) -> list[object]:
            raise RuntimeError("boom")

    # The except-Exception boundary must swallow the error, not propagate it.
    _normalize_table_widths(Boom())  # must not raise
